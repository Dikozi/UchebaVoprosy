#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Планер привычек — генератор Excel-файла.

Собирает готовый .xlsx:
  • «Старт»  — год, список привычек (сфера, название, цель дней в месяц, зачем), инструкция;
  • 12 листов-месяцев — сетка «привычки × дни» с нативными чекбоксами Excel 365,
    статистика по дням / неделям / привычкам, серии 🔥, KPI-плитки, инсайты,
    графики, строки сна / энергии / настроения и дневник месяца;
  • «Год»    — итоги по месяцам, график, тепловая карта «привычки × месяцы»;
  • «Списки» — скрытый лист со справочниками для выпадающих списков.

Запуск:
    python3 build_planner.py                 # соберёт файлы в ./dist
    python3 build_planner.py --out DIR       # другая папка
    python3 build_planner.py --recalc        # двухпроходная сборка: формулы пересчитываются
                                             # в LibreOffice (headless), проверяются на ошибки,
                                             # а значения записываются в файл как кэш — цифры
                                             # видны даже в превью, где формулы не считаются.
    python3 build_planner.py --preview       # дополнительно собрать вариант для рендера
                                             # (галочки как «✓» вместо чекбоксов)
"""
import argparse
import calendar
import datetime as dt
import os
import random
import subprocess

import xlsxwriter
from xlsxwriter.utility import xl_col_to_name as cn
from xlsxwriter.utility import xl_rowcol_to_cell as rc

# ----------------------------------------------------------------------------
# Справочники
# ----------------------------------------------------------------------------
MONTHS = [("Январь", "янв"), ("Февраль", "фев"), ("Март", "мар"), ("Апрель", "апр"),
          ("Май", "мая"), ("Июнь", "июн"), ("Июль", "июл"), ("Август", "авг"),
          ("Сентябрь", "сен"), ("Октябрь", "окт"), ("Ноябрь", "ноя"), ("Декабрь", "дек")]
MAX_DAYS = [31, 29, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]   # максимум дней в месяце
WD_LIST = '"Пн","Вт","Ср","Чт","Пт","Сб","Вс"'
MOODS = ["😄", "🙂", "😐", "😕", "😫"]
CATEGORIES = [("Здоровье", "#34D399"), ("Спорт", "#38BDF8"), ("Работа", "#F59E0B"),
              ("Учёба", "#A78BFA"), ("Личное", "#F472B6"), ("Финансы", "#FCD34D"),
              ("Отношения", "#FB7185"), ("Творчество", "#22D3EE")]
QUOTES = [
    "Ты — это то, что ты делаешь каждый день.",
    "Маленькие шаги каждый день и есть большой результат.",
    "Мотивация запускает. Привычка везёт.",
    "Не идеально, а регулярно.",
    "Дисциплина — выбор между «хочу сейчас» и «хочу больше всего».",
    "Сегодня — самый важный день серии.",
    "Плюс 1% в день — в 37 раз лучше за год.",
    "Пропустить один раз — случайность. Дважды — новая привычка.",
    "Система важнее цели.",
    "Прогресс, а не совершенство.",
    "Каждая галочка — голос за того, кем ты становишься.",
    "Год — это 365 маленьких решений.",
]
SEASON = {1: "#38BDF8", 2: "#38BDF8", 3: "#34D399", 4: "#34D399", 5: "#34D399",
          6: "#FB923C", 7: "#FB923C", 8: "#FB923C", 9: "#F59E0B", 10: "#F59E0B",
          11: "#F59E0B", 12: "#38BDF8"}
EXAMPLE_HABITS = [
    ("Здоровье", "Подъём до 7:00", None, "Утро задаёт весь день"),
    ("Спорт", "Тренировка", 12, "Сила, энергия, форма"),
    ("Здоровье", "10 000 шагов", None, "Голова яснее, сон крепче"),
    ("Учёба", "Чтение 20 минут", None, "24 книги за год"),
    ("Личное", "Медитация 10 минут", None, "Спокойствие и фокус"),
    ("Здоровье", "Без сахара", None, "Ровная энергия весь день"),
    ("Работа", "План на завтра", None, "Утро без хаоса"),
    ("Учёба", "Английский 15 минут", 20, "Свободно говорить к лету"),
]
N_HABITS = 15
YEAR_REF = "'Старт'!$D$6"
LISTS = "'Списки'"
FONT = "Arial"

# ----------------------------------------------------------------------------
# Темы оформления
# ----------------------------------------------------------------------------
THEMES = {
    "dark": dict(
        bg="#0E1016", panel="#171A23", panel2="#20242F", tile="#171A23", input="#2A2617",
        line="#2A2F3D", line2="#3A4052", text="#ECEEF3", muted="#8A90A3", dim="#5E6478",
        accent="#F6C453", accent_text="#1A1400", green="#2FD07A", green_text="#F3FFF7",
        red="#FF5C7A", weekend="#1B1E29", today="#3B3320",
        weeks=["#E0456A", "#7C4DE8", "#2A9FD6", "#E8792F", "#22B183"],
        heat=("#7F1D1D", "#8A6A12", "#166534"), chart_bg="#171A23", grid="#2A2F3D",
        pill_mix=0.80, pill_font=0.0,
    ),
    "light": dict(
        bg="#FFFFFF", panel="#F5F6FA", panel2="#ECEEF5", tile="#F5F6FA", input="#FFF6D6",
        line="#DDE1EA", line2="#C9CFDC", text="#171A24", muted="#6B7280", dim="#9AA1B2",
        accent="#F2B300", accent_text="#1A1400", green="#22C55E", green_text="#FFFFFF",
        red="#E11D48", weekend="#EEF0F6", today="#FFF1BF",
        weeks=["#E11D48", "#7C3AED", "#0284C7", "#EA580C", "#059669"],
        heat=("#FCA5A5", "#FDE68A", "#86EFAC"), chart_bg="#F5F6FA", grid="#E3E6EE",
        pill_mix=0.82, pill_font=0.5,
    ),
}


def mix(c1, c2, t):
    """Смешать два цвета: t=0 → c1, t=1 → c2."""
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02X%02X%02X" % tuple(round(x * (1 - t) + y * t) for x, y in zip(a, b))


# ----------------------------------------------------------------------------
# Геометрия листа месяца (0-based индексы)
# ----------------------------------------------------------------------------
C_MARGIN, C_CAT, C_HABIT, C_TARGET, C_D1 = 0, 1, 2, 3, 4
C_D31 = C_D1 + 30
C_SP = C_D31 + 1
C_DONE, C_PROG, C_STREAK, C_BEST = C_SP + 1, C_SP + 2, C_SP + 3, C_SP + 4
C_SP2 = C_BEST + 1
C_H = C_SP2 + 1            # скрытая колонка служебных значений
C_RUN = C_H + 1            # скрытый блок серий (31 колонка)
C_LAST = C_RUN + 30

R_TITLE, R_SUB = 1, 2
R_KL, R_KV, R_KS = 4, 5, 6                       # KPI-плитки: подпись / значение / пояснение
R_BAND, R_WD, R_DAY = 8, 9, 10                   # недели / дни недели / числа
R_H0 = 11
R_H1 = R_H0 + N_HABITS - 1                       # 25
R_CNT, R_PCT = R_H1 + 1, R_H1 + 2                # 26, 27
R_SLEEP, R_ENERGY, R_MOOD = R_H1 + 4, R_H1 + 5, R_H1 + 6   # 29, 30, 31
R_SEC = R_MOOD + 1                               # 32 — заголовки секций графиков
R_CHART0 = R_SEC + 1                             # 33
R_CHART1 = R_CHART0 + 13                         # 46
R_DT = R_CHART1 + 2                              # 48 — заголовок дневника
R_DH = R_DT + 1                                  # 49
R_D0 = R_DH + 1                                  # 50 — первый день дневника
R_D31 = R_D0 + 30                                # 80
R_MOODCNT = 29                                   # счётчики настроений в скрытой колонке


def H(i):
    return rc(i, C_H, True, True)


FIRST, DIM, EL, NH, PLAN, DONE, PROG, REST, PERF = (H(i) for i in range(9))


def A(r, c, ra=False, ca=False):
    return rc(r, c, ra, ca)


def R(r1, c1, r2, c2):
    return f"{rc(r1, c1, True, True)}:{rc(r2, c2, True, True)}"


NAMES = R(R_H0, C_HABIT, R_H1, C_HABIT)
DAY_HDR = R(R_DAY, C_D1, R_DAY, C_D31)
CNT_RNG = R(R_CNT, C_D1, R_CNT, C_D31)
PCT_RNG = R(R_PCT, C_D1, R_PCT, C_D31)
SLEEP_RNG = R(R_SLEEP, C_D1, R_SLEEP, C_D31)
ENERGY_RNG = R(R_ENERGY, C_D1, R_ENERGY, C_D31)
MOOD_RNG = R(R_MOOD, C_D1, R_MOOD, C_D31)
PROG_RNG = R(R_H0, C_PROG, R_H1, C_PROG)
BEST_RNG = R(R_H0, C_BEST, R_H1, C_BEST)
MOODCNT_RNG = R(R_MOODCNT, C_H, R_MOODCNT + len(MOODS) - 1, C_H)

EMPTY = dict(checks=set(), sleep={}, energy={}, mood={}, diary={})


# ----------------------------------------------------------------------------
# Контекст сборки: кэш форматов, запись формул с кэшем значений, чекбоксы
# ----------------------------------------------------------------------------
class Ctx:
    def __init__(self, wb, theme, cache=None, preview=False):
        self.wb, self.th, self.cache, self.preview = wb, theme, cache, preview
        self._fmts = {}
        self.formula_cells = set()

    def F(self, **kw):
        checkbox = kw.pop("_cb", False)
        key = tuple(sorted(kw.items())) + (("_cb", checkbox),)
        if key not in self._fmts:
            fmt = self.wb.add_format(kw)
            if checkbox:
                fmt.set_checkbox()
            self._fmts[key] = fmt
        return self._fmts[key]

    def fx(self, ws, row, col, formula, fmt=None):
        """Формула + (при втором проходе) её кэшированное значение."""
        self.formula_cells.add((ws.name, row, col))
        value = 0
        if self.cache is not None:
            value = self.cache.get((ws.name, row, col))
            if value is None:
                value = ""
        ws.write_formula(row, col, formula, fmt, value)

    def cached(self, ws, row, col):
        return None if self.cache is None else self.cache.get((ws.name, row, col))

    def cb(self, ws, row, col, checked, fmt, preview_fmt):
        """Чекбокс Excel 365. В режиме превью — число 1 с форматом «✓» (для рендера)."""
        if self.preview:
            if checked:
                ws.write_number(row, col, 1, preview_fmt)
            else:
                ws.write_blank(row, col, None, preview_fmt)
        else:
            ws.insert_checkbox(row, col, bool(checked), fmt)


# ----------------------------------------------------------------------------
# Лист месяца
# ----------------------------------------------------------------------------
def build_month(ctx, mi, mdata, today):
    th, wb = ctx.th, ctx.wb
    name, short = MONTHS[mi - 1]
    ws = wb.add_worksheet(name)
    S = f"'{name}'"
    ws.hide_gridlines(2)
    ws.set_zoom(90)
    ws.set_tab_color(SEASON[mi])
    ws.set_default_row(15)

    base = ctx.F(bg_color=th["bg"], font_color=th["text"])
    hidden = ctx.F(bg_color=th["bg"], font_color=th["dim"])

    # --- колонки ---
    ws.set_column(C_LAST + 1, 16383, None, base)
    ws.set_column(C_MARGIN, C_MARGIN, 1.5, base)
    ws.set_column(C_CAT, C_CAT, 12, base)
    ws.set_column(C_HABIT, C_HABIT, 26, base)
    ws.set_column(C_TARGET, C_TARGET, 6.5, base)
    ws.set_column(C_D1, C_D31, 3.5, base)
    ws.set_column(C_SP, C_SP, 1.5, base)
    ws.set_column(C_DONE, C_DONE, 8.5, base)
    ws.set_column(C_PROG, C_PROG, 10, base)
    ws.set_column(C_STREAK, C_STREAK, 7.5, base)
    ws.set_column(C_BEST, C_BEST, 7.5, base)
    ws.set_column(C_SP2, C_SP2, 1.5, base)
    ws.set_column(C_H, C_LAST, 4, hidden, {"hidden": True})

    # --- высоты строк ---
    heights = {0: 6, R_TITLE: 30, R_SUB: 16, 3: 8, R_KL: 14, R_KV: 28, R_KS: 14, 7: 8,
               R_BAND: 18, R_WD: 13, R_DAY: 18, R_CNT: 16, R_PCT: 16, R_H1 + 3: 6,
               R_SLEEP: 18, R_ENERGY: 18, R_MOOD: 18, R_SEC: 18, R_CHART1 + 1: 8,
               R_DT: 22, R_DH: 16}
    for r in range(R_H0, R_H1 + 1):
        heights[r] = 20
    for r in range(R_CHART0, R_CHART1 + 1):
        heights[r] = 15
    for r in range(R_D0, R_D31 + 1):
        heights[r] = 17
    for r, h in heights.items():
        ws.set_row(r, h)

    # --- служебные значения (скрытая колонка) ---
    ctx.fx(ws, 0, C_H, f"=DATE({YEAR_REF},{mi},1)", hidden)
    ctx.fx(ws, 1, C_H, f"=DAY(EOMONTH({FIRST},0))", hidden)
    ctx.fx(ws, 2, C_H, f"=IF(TODAY()<{FIRST},0,IF(TODAY()>EOMONTH({FIRST},0),{DIM},DAY(TODAY())))", hidden)
    ctx.fx(ws, 3, C_H, f'=SUMPRODUCT(--({NAMES}<>""))', hidden)
    ctx.fx(ws, 4, C_H, f"=ROUND(SUM({R(R_H0, C_H, R_H1, C_H)}),0)", hidden)
    ctx.fx(ws, 5, C_H, f"=SUM({R(R_H0, C_DONE, R_H1, C_DONE)})", hidden)
    ctx.fx(ws, 6, C_H, f"=IFERROR(AVERAGE({PROG_RNG}),0)",
           ctx.F(bg_color=th["bg"], font_color=th["dim"], num_format="0%"))
    ctx.fx(ws, 7, C_H, f"=1-{PROG}", hidden)
    ctx.fx(ws, 8, C_H, f"=SUMPRODUCT(--({PCT_RNG}=1))", hidden)
    for i in range(len(MOODS)):
        ctx.fx(ws, R_MOODCNT + i, C_H, f"=COUNTIF({MOOD_RNG},{LISTS}!$C${2 + i})", hidden)

    # --- заголовок ---
    tf = ctx.F(bg_color=th["bg"], font_color=th["text"], bold=True, font_size=22, valign="vcenter")
    ws.merge_range(R_TITLE, C_CAT, R_TITLE, C_D1 + 12, "", tf)
    ctx.fx(ws, R_TITLE, C_CAT, f'="{name.upper()} "&{YEAR_REF}', tf)
    brand = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=8, bold=True,
                  align="right", valign="vcenter")
    ws.merge_range(R_TITLE, C_D1 + 13, R_TITLE, C_BEST, "ПЛАНЕР ПРИВЫЧЕК", brand)
    sub = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=9, valign="vcenter")
    ws.merge_range(R_SUB, C_CAT, R_SUB, C_D1 + 12,
                   "Ставь галочку каждый день · сон, энергия и настроение — под сеткой · "
                   "список привычек и цели — на листе «Старт»", sub)
    quote = ctx.F(bg_color=th["bg"], font_color=th["accent"], font_size=9, italic=True,
                  align="right", valign="vcenter")
    ws.merge_range(R_SUB, C_D1 + 13, R_SUB, C_BEST, QUOTES[mi - 1], quote)

    # --- KPI-плитки ---
    tile_kw = dict(bg_color=th["tile"], left=5, left_color=th["accent"], align="left", indent=1)
    kl = ctx.F(font_color=th["muted"], font_size=8, bold=True, valign="bottom", **tile_kw)
    kv = ctx.F(font_color=th["text"], font_size=18, bold=True, valign="vcenter", **tile_kw)
    kv_pct = ctx.F(font_color=th["accent"], font_size=18, bold=True, valign="vcenter",
                   num_format="0%", **tile_kw)
    kv_days = ctx.F(font_color=th["text"], font_size=18, bold=True, valign="vcenter",
                    num_format='0" дн."', **tile_kw)
    kv_txt = ctx.F(font_color=th["text"], font_size=14, bold=True, valign="vcenter", **tile_kw)
    ks = ctx.F(font_color=th["muted"], font_size=8, valign="top", **tile_kw)
    started = f"{EL}*{NH}*{DONE}=0"          # месяц не начался, нет привычек или нет ни одной отметки
    tiles = [
        ("ПРОГРЕСС МЕСЯЦА", f"={PROG}", kv_pct,
         f'=IF({EL}=0,"месяц ещё не начался",{DONE}&" из "&{PLAN}&" отметок · день "&{EL}&" из "&{DIM})'),
        ("ИДЕАЛЬНЫХ ДНЕЙ", f"={PERF}", kv,
         f'=IF({EL}=0,"—","из "&{EL}&" прошедших · все привычки ✓")'),
        ("ЛУЧШАЯ СЕРИЯ 🔥", f"=IFERROR(MAX({BEST_RNG}),0)", kv_days,
         f'=IF(IFERROR(MAX({BEST_RNG}),0)=0,"серий пока нет",'
         f'INDEX({NAMES},MATCH(MAX({BEST_RNG}),{BEST_RNG},0)))'),
        ("ЛУЧШИЙ ДЕНЬ", f'=IF({started},"—",INDEX({DAY_HDR},MATCH(MAX({PCT_RNG}),{PCT_RNG},0))&" {short}")',
         kv_txt, f'=IF({started},"",ROUND(MAX({PCT_RNG})*100,0)&"% привычек выполнено")'),
        ("СЛАБЫЙ ДЕНЬ", f'=IF({started},"—",INDEX({DAY_HDR},MATCH(MIN({PCT_RNG}),{PCT_RNG},0))&" {short}")',
         kv_txt, f'=IF({started},"",ROUND(MIN({PCT_RNG})*100,0)&"% привычек выполнено")'),
        ("ТОП ПРИВЫЧКА", f'=IF({started},"—",INDEX({NAMES},MATCH(MAX({PROG_RNG}),{PROG_RNG},0)))',
         kv_txt, f'=IF({started},"","Подтянуть: "&INDEX({NAMES},MATCH(MIN({PROG_RNG}),{PROG_RNG},0)))'),
    ]
    spans = [(C_CAT, C_TARGET), (C_D1, C_D1 + 6), (C_D1 + 7, C_D1 + 13), (C_D1 + 14, C_D1 + 20),
             (C_D1 + 21, C_D1 + 27), (C_D1 + 28, C_BEST)]
    for (label, fval, vfmt, fsub), (c1, c2) in zip(tiles, spans):
        ws.merge_range(R_KL, c1, R_KL, c2, label, kl)
        ws.merge_range(R_KV, c1, R_KV, c2, "", vfmt)
        ctx.fx(ws, R_KV, c1, fval, vfmt)
        ws.merge_range(R_KS, c1, R_KS, c2, "", ks)
        ctx.fx(ws, R_KS, c1, fsub, ks)

    # --- полоса недель + заголовки секций ---
    sec = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter")
    ws.merge_range(R_BAND, C_CAT, R_BAND, C_TARGET, "МОИ ПРИВЫЧКИ", sec)
    ws.merge_range(R_BAND, C_DONE, R_BAND, C_BEST, "СТАТИСТИКА", sec)
    for w in range(5):
        a, b = 7 * w + 1, min(7 * w + 7, 31)
        ca, cb_ = C_D1 + a - 1, C_D1 + b - 1
        plan = f"({NH}*MAX(0,MIN({EL},{b})-{a}+1))"
        done = f"SUM({R(R_CNT, ca, R_CNT, cb_)})"
        f = (f'=IF({a}>{DIM},"",IF({plan}=0,"НЕДЕЛЯ {w + 1}","НЕДЕЛЯ {w + 1} · "&{done}&"/"&{plan}'
             f'&" · "&ROUND({done}/{plan}*100,0)&"%"))')
        bf = ctx.F(bg_color=th["weeks"][w], font_color="#FFFFFF", bold=True, font_size=8,
                   align="center", valign="vcenter")
        ws.merge_range(R_BAND, ca, R_BAND, cb_, "", bf)
        ctx.fx(ws, R_BAND, ca, f, bf)
    ws.conditional_format(R_BAND, C_D1 + 28, R_BAND, C_D31, {
        "type": "formula", "criteria": f'={A(R_BAND, C_D1 + 28)}=""',
        "format": ctx.F(bg_color=th["bg"], font_color=th["bg"])})

    # --- шапка сетки ---
    hdr = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True,
                valign="vcenter", align="left", indent=1, bottom=1, bottom_color=th["line2"])
    hdr_c = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True,
                  valign="vcenter", align="center", bottom=1, bottom_color=th["line2"])
    for c, text in ((C_CAT, "Сфера"), (C_HABIT, "Привычка")):
        ws.merge_range(R_WD, c, R_DAY, c, text, hdr)
    ws.merge_range(R_WD, C_TARGET, R_DAY, C_TARGET, "Цель", hdr_c)
    for c, text in ((C_DONE, "Сделано"), (C_PROG, "Прогресс"), (C_STREAK, "Серия 🔥"), (C_BEST, "Рекорд")):
        ws.merge_range(R_WD, c, R_DAY, c, text, hdr_c)
    ws.write_comment(R_WD,C_TARGET,
                     "Цель — сколько дней в месяц планируешь выполнять привычку. "
                     "Пусто = каждый день. Задаётся на листе «Старт», можно перезаписать здесь.",
                     {"x_scale": 1.6, "y_scale": 1.1})
    ws.write_comment(R_WD,C_HABIT,
                     "Список берётся с листа «Старт». Хочешь другой набор только на этот месяц — "
                     "просто впиши название прямо сюда.", {"x_scale": 1.6, "y_scale": 1.0})
    ws.write_comment(R_WD,C_PROG,
                     "Прогресс = сделано / ожидаемо к сегодняшнему дню. Для ежедневной привычки — "
                     "доля отмеченных дней; для привычки с целью — доля от цели (пропорционально дню месяца).",
                     {"x_scale": 1.8, "y_scale": 1.2})
    ws.write_comment(R_WD,C_STREAK,
                     "Серия — сколько дней подряд выполнена привычка (сегодняшний день не ломает серию, "
                     "пока не закончился). Рекорд — самая длинная серия месяца.",
                     {"x_scale": 1.8, "y_scale": 1.1})
    hdr_wd = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=7, align="center", valign="vcenter")
    hdr_day = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=9, bold=True,
                    align="center", valign="vcenter", bottom=1, bottom_color=th["line2"])
    for k in range(1, 32):
        c = C_D1 + k - 1
        ctx.fx(ws, R_WD, c, f'=IF({k}<={DIM},CHOOSE(WEEKDAY({FIRST}+{k - 1},2),{WD_LIST}),"")', hdr_wd)
        ctx.fx(ws, R_DAY, c, f'=IF({k}<={DIM},{k},"")', hdr_day)

    # --- строки привычек ---
    cat_f = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, align="center", valign="vcenter")
    name_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=10, bold=True,
                   valign="vcenter", align="left", indent=1)
    tgt_f = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=9, align="center", valign="vcenter")
    cb_f = ctx.F(bg_color=th["panel2"], font_color=th["dim"], font_size=9, align="center",
                 valign="vcenter", border=1, border_color=th["bg"], _cb=True)
    cb_prev = ctx.F(bg_color=th["panel2"], font_color=th["dim"], font_size=9, align="center",
                    valign="vcenter", border=1, border_color=th["bg"], num_format='"✓"')
    void_f = ctx.F(bg_color=th["bg"])
    done_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=10, bold=True,
                   align="center", valign="vcenter")
    prog_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=9, align="center",
                   valign="vcenter", num_format="0%")
    streak_f = ctx.F(bg_color=th["panel"], font_color=th["accent"], font_size=10, bold=True,
                     align="center", valign="vcenter", num_format='[>=3]0" 🔥";0')
    best_f = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=9, align="center", valign="vcenter")
    hc, nc, tc = cn(C_CAT), cn(C_HABIT), cn(C_TARGET)
    for i in range(N_HABITS):
        r = R_H0 + i
        sr = r + 1                       # та же строка на листе «Старт»
        ctx.fx(ws, r, C_CAT, f"=IF('Старт'!$C${sr}=\"\",\"\",'Старт'!$C${sr})", cat_f)
        ctx.fx(ws, r, C_HABIT, f"=IF('Старт'!$D${sr}=\"\",\"\",'Старт'!$D${sr})", name_f)
        ctx.fx(ws, r, C_TARGET, f"=IF('Старт'!$E${sr}=\"\",\"\",'Старт'!$E${sr})", tgt_f)
        for k in range(1, 32):
            c = C_D1 + k - 1
            if k > MAX_DAYS[mi - 1]:
                ws.write_blank(r, c, None, void_f)
            else:
                ctx.cb(ws, r, c, (i, k) in mdata["checks"], cb_f, cb_prev)
        row_rng = f"{A(r, C_D1)}:{A(r, C_D31)}"
        run_rng = f"${cn(C_RUN)}${sr}:${cn(C_LAST)}${sr}"
        ctx.fx(ws, r, C_DONE,
               f'=IF(${nc}{sr}="","",SUMPRODUCT(((({row_rng}=TRUE)+({row_rng}=1))>0)*({DAY_HDR}<>"")))', done_f)
        ctx.fx(ws, r, C_PROG,
               f'=IF(${nc}{sr}="","",IF({EL}=0,0,IFERROR(MIN(1,{A(r, C_DONE)}/{H(r)}),0)))', prog_f)
        ctx.fx(ws, r, C_STREAK,
               f'=IF(${nc}{sr}="","",IF({EL}=0,0,MAX(INDEX({run_rng},{EL}),INDEX({run_rng},MAX({EL}-1,1)))))',
               streak_f)
        ctx.fx(ws, r, C_BEST, f'=IF(${nc}{sr}="","",MAX({run_rng}))', best_f)
        ctx.fx(ws, r, C_H, f'=IF(${nc}{sr}="",0,IFERROR(IF(${tc}{sr}="",{DIM},${tc}{sr})*{EL}/{DIM},{EL}))', hidden)
        for k in range(1, 32):
            c = C_RUN + k - 1
            dc = A(r, C_D1 + k - 1)
            prev = "0" if k == 1 else A(r, c - 1)
            ctx.fx(ws, r, c, f"=IF(AND({k}<={EL},OR({dc}=TRUE,{dc}=1)),{prev}+1,0)", hidden)

    # --- итоги по дням ---
    lab = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=8, align="right", valign="vcenter")
    cnt_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=8, bold=True, align="center", valign="vcenter")
    pct_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=7, align="center",
                  valign="vcenter", num_format="0%")
    ws.merge_range(R_CNT, C_CAT, R_CNT, C_TARGET, "Выполнено за день", lab)
    ws.merge_range(R_PCT, C_CAT, R_PCT, C_TARGET, "Процент дня", lab)
    for k in range(1, 32):
        c = C_D1 + k - 1
        col = cn(c)
        colrng = f"{col}${R_H0 + 1}:{col}${R_H1 + 1}"
        ctx.fx(ws, R_CNT, c,
               f'=IF(OR({col}${R_DAY + 1}="",{col}${R_DAY + 1}>{EL}),"",'
               f'SUMPRODUCT(((({colrng}=TRUE)+({colrng}=1))>0)*({NAMES}<>"")))', cnt_f)
        ctx.fx(ws, R_PCT, c, f'=IF(OR({col}{R_CNT + 1}="",{NH}=0),"",{col}{R_CNT + 1}/{NH})', pct_f)
    tot_f = ctx.F(bg_color=th["panel"], font_color=th["accent"], font_size=9, bold=True, align="center", valign="vcenter")
    tot_p = ctx.F(bg_color=th["panel"], font_color=th["accent"], font_size=9, bold=True, align="center",
                  valign="vcenter", num_format="0%")
    ctx.fx(ws, R_CNT, C_DONE, f"={DONE}", tot_f)
    ws.merge_range(R_CNT, C_PROG, R_CNT, C_BEST, "всего отметок", ctx.F(
        bg_color=th["panel"], font_color=th["muted"], font_size=8, valign="vcenter", align="left", indent=1))
    ctx.fx(ws, R_PCT, C_DONE, f"={PROG}", tot_p)
    ws.merge_range(R_PCT, C_PROG, R_PCT, C_BEST, "прогресс месяца", ctx.F(
        bg_color=th["panel"], font_color=th["muted"], font_size=8, valign="vcenter", align="left", indent=1))

    # --- сон / энергия / настроение ---
    in_f = ctx.F(bg_color=th["input"], font_color=th["text"], font_size=8, align="center",
                 valign="vcenter", border=1, border_color=th["bg"])
    ws.merge_range(R_SLEEP, C_CAT, R_SLEEP, C_TARGET, "Сон, часов", lab)
    ws.merge_range(R_ENERGY, C_CAT, R_ENERGY, C_TARGET, "Энергия, 1–5", lab)
    ws.merge_range(R_MOOD, C_CAT, R_MOOD, C_TARGET, "Настроение", lab)
    for k in range(1, 32):
        c = C_D1 + k - 1
        for r, key in ((R_SLEEP, "sleep"), (R_ENERGY, "energy"), (R_MOOD, "mood")):
            v = mdata[key].get(k)
            if v is None:
                ws.write_blank(r, c, None, in_f)
            elif isinstance(v, str):
                ws.write_string(r, c, v, in_f)
            else:
                ws.write_number(r, c, v, in_f)
    ws.data_validation(R_SLEEP, C_D1, R_SLEEP, C_D31, {
        "validate": "decimal", "criteria": "between", "minimum": 0, "maximum": 24, "ignore_blank": True,
        "error_title": "Сон", "error_message": "Введи часы сна числом от 0 до 24, например 7,5"})
    ws.data_validation(R_ENERGY, C_D1, R_ENERGY, C_D31, {
        "validate": "list", "source": f"={LISTS}!$B$2:$B$6", "ignore_blank": True,
        "error_title": "Энергия", "error_message": "Оцени энергию от 1 (нет сил) до 5 (полон сил)"})
    ws.data_validation(R_MOOD, C_D1, R_MOOD, C_D31, {
        "validate": "list", "source": f"={LISTS}!$C$2:$C$6", "ignore_blank": True,
        "error_title": "Настроение", "error_message": "Выбери смайлик из списка"})
    hint = ctx.F(bg_color=th["bg"], font_color=th["dim"], font_size=7, valign="vcenter", align="left", indent=1)
    ws.merge_range(R_SLEEP, C_DONE, R_SLEEP, C_BEST, "часы, например 7,5", hint)
    ws.merge_range(R_ENERGY, C_DONE, R_ENERGY, C_BEST, "1 — нет сил, 5 — полон сил", hint)
    ws.merge_range(R_MOOD, C_DONE, R_MOOD, C_BEST, "выбери смайлик из списка", hint)

    # --- условное форматирование сетки ---
    e0 = A(R_H0, C_D1)                                   # верхняя левая ячейка сетки (относительная)
    day0 = f"{cn(C_D1)}${R_DAY + 1}"                     # число дня для колонки
    wd0 = f"{cn(C_D1)}${R_WD + 1}"
    checked = ctx.F(bg_color=th["green"], font_color=th["green_text"], bold=True,
                    border=1, border_color=th["bg"])
    beyond = ctx.F(bg_color=th["bg"], font_color=th["bg"], border=1, border_color=th["bg"])
    today_f = ctx.F(bg_color=th["today"], font_color=th["accent"], border=1, border_color=th["bg"])
    weekend_f = ctx.F(bg_color=th["weekend"], border=1, border_color=th["bg"])
    ws.conditional_format(R_H0, C_D1, R_H1, C_D31, {
        "type": "formula", "criteria": f"=OR({e0}=TRUE,{e0}=1)", "format": checked, "stop_if_true": True})
    grid_rows = [(R_H0, R_H1), (R_SLEEP, R_MOOD)]
    for r1, r2 in grid_rows:
        ws.conditional_format(r1, C_D1, r2, C_D31, {
            "type": "formula", "criteria": f'={day0}=""', "format": beyond, "stop_if_true": True})
        ws.conditional_format(r1, C_D1, r2, C_D31, {
            "type": "formula", "criteria": f"=AND({day0}<>\"\",{FIRST}+{day0}-1=TODAY())",
            "format": today_f, "stop_if_true": True})
        ws.conditional_format(r1, C_D1, r2, C_D31, {
            "type": "formula", "criteria": f'=OR({wd0}="Сб",{wd0}="Вс")', "format": weekend_f})
    # сон / энергия: цветовые подсказки
    s0 = A(R_SLEEP, C_D1)
    en0 = A(R_ENERGY, C_D1)
    ws.conditional_format(R_SLEEP, C_D1, R_SLEEP, C_D31, {
        "type": "formula", "criteria": f"=AND(ISNUMBER({s0}),{s0}<6.5)", "format": ctx.F(font_color=th["red"], bold=True)})
    ws.conditional_format(R_SLEEP, C_D1, R_SLEEP, C_D31, {
        "type": "formula", "criteria": f"=AND(ISNUMBER({s0}),{s0}>=7.5)", "format": ctx.F(font_color=th["green"], bold=True)})
    ws.conditional_format(R_ENERGY, C_D1, R_ENERGY, C_D31, {
        "type": "formula", "criteria": f"=AND(ISNUMBER({en0}),{en0}<=2)", "format": ctx.F(font_color=th["red"], bold=True)})
    ws.conditional_format(R_ENERGY, C_D1, R_ENERGY, C_D31, {
        "type": "formula", "criteria": f"=AND(ISNUMBER({en0}),{en0}>=4)", "format": ctx.F(font_color=th["green"], bold=True)})
    # шапка: сегодня / выходные / вне месяца
    hdr_today = ctx.F(bg_color=th["accent"], font_color=th["accent_text"], bold=True)
    hdr_beyond = ctx.F(bg_color=th["bg"], font_color=th["bg"])
    hdr_weekend = ctx.F(font_color=th["red"])
    for r in (R_WD, R_DAY):
        ws.conditional_format(r, C_D1, r, C_D31, {
            "type": "formula", "criteria": f'={day0}=""', "format": hdr_beyond, "stop_if_true": True})
        ws.conditional_format(r, C_D1, r, C_D31, {
            "type": "formula", "criteria": f"=AND({day0}<>\"\",{FIRST}+{day0}-1=TODAY())",
            "format": hdr_today, "stop_if_true": True})
        ws.conditional_format(r, C_D1, r, C_D31, {
            "type": "formula", "criteria": f'=OR({wd0}="Сб",{wd0}="Вс")', "format": hdr_weekend})
    # итоги по дням: тепловая шкала и скрытие пустых
    for r in (R_CNT, R_PCT):
        ws.conditional_format(r, C_D1, r, C_D31, {
            "type": "formula", "criteria": f'={day0}=""', "format": hdr_beyond, "stop_if_true": True})
    ws.conditional_format(R_PCT, C_D1, R_PCT, C_D31, {
        "type": "3_color_scale", "min_type": "num", "min_value": 0, "mid_type": "num", "mid_value": 0.5,
        "max_type": "num", "max_value": 1, "min_color": th["heat"][0], "mid_color": th["heat"][1],
        "max_color": th["heat"][2]})
    # прогресс привычки: полоса
    ws.conditional_format(R_H0, C_PROG, R_H1, C_PROG, {
        "type": "data_bar", "bar_color": th["green"], "bar_solid": True, "bar_no_border": True,
        "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 1, "bar_axis_position": "none"})
    # сферы: цветные «пилюли»
    add_category_pills(ctx, ws, R_H0, C_CAT, R_H1, C_CAT)

    # --- графики + инсайты ---
    ws.merge_range(R_SEC, C_CAT, R_SEC, C_TARGET, "ПРОГРЕСС МЕСЯЦА", sec)
    ws.merge_range(R_SEC, C_D1, R_SEC, C_D31, "ВЫПОЛНЕНИЕ ПО ДНЯМ", sec)
    ws.merge_range(R_SEC, C_DONE, R_SEC, C_BEST, "ИНСАЙТЫ МЕСЯЦА", sec)
    chart_h = 14 * 20
    donut = wb.add_chart({"type": "doughnut"})
    donut.add_series({
        "values": f"={S}!{PROG}:{REST}",
        "points": [{"fill": {"color": th["green"]}, "border": {"none": True}},
                   {"fill": {"color": th["panel2"]}, "border": {"none": True}}],
    })
    donut.set_hole_size(72)
    donut.set_legend({"none": True})
    donut.set_title({"none": True})
    donut.set_chartarea({"fill": {"color": th["chart_bg"]}, "border": {"none": True}})
    donut.set_plotarea({"fill": {"none": True}})
    donut_w = 12 * 7 + 5 + 26 * 7 + 5 + int(6.5 * 7 + 5)
    donut.set_size({"width": donut_w, "height": chart_h})
    ws.insert_chart(R_CHART0, C_CAT, donut)
    pv = ctx.cached(ws, 6, C_H)
    center = f"{int(pv * 100 + 0.5)}%" if isinstance(pv, (int, float)) else ""
    ws.insert_textbox(R_CHART0, C_CAT, center, {
        "width": donut_w, "height": chart_h, "textlink": f"={S}!{PROG}",
        "font": {"name": FONT, "size": 24, "bold": True, "color": th["text"]},
        "align": {"vertical": "middle", "horizontal": "center"},
        "fill": {"none": True}, "line": {"none": True}})

    col = wb.add_chart({"type": "column"})
    col.add_series({
        "name": "Процент дня", "categories": f"={S}!{DAY_HDR}", "values": f"={S}!{PCT_RNG}",
        "fill": {"color": th["accent"]}, "border": {"none": True}, "gap": 55,
    })
    col.set_legend({"none": True})
    col.set_title({"none": True})
    col.set_chartarea({"fill": {"color": th["chart_bg"]}, "border": {"none": True}})
    col.set_plotarea({"fill": {"none": True}})
    axis_font = {"name": FONT, "color": th["muted"], "size": 8}
    col.set_x_axis({"num_font": axis_font, "line": {"color": th["grid"]},
                    "major_gridlines": {"visible": False}, "major_tick_mark": "none"})
    col.set_y_axis({"min": 0, "max": 1, "major_unit": 0.25, "num_format": "0%", "num_font": axis_font,
                    "line": {"none": True}, "major_gridlines": {"visible": True, "line": {"color": th["grid"]}}})
    col.set_size({"width": 31 * int(3.5 * 7 + 5), "height": chart_h})
    ws.insert_chart(R_CHART0, C_D1, col)

    il = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=7, valign="bottom", align="left", indent=1)
    iv = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=11, bold=True, valign="top", align="left", indent=1)
    insights = [
        ("Средний сон", f'=IFERROR(ROUND(AVERAGE({SLEEP_RNG}),1)&" ч","—")'),
        ("Средняя энергия", f'=IFERROR(ROUND(AVERAGE({ENERGY_RNG}),1)&" из 5","—")'),
        ("Частое настроение", f'=IF(MAX({MOODCNT_RNG})=0,"—",INDEX({LISTS}!$C$2:$C$6,MATCH(MAX({MOODCNT_RNG}),{MOODCNT_RNG},0)))'),
        ("Прогресс при сне ≥ 7 ч", f'=IFERROR(ROUND(AVERAGEIFS({PCT_RNG},{SLEEP_RNG},">=7")*100,0)&"%","—")'),
        ("Прогресс при сне < 7 ч", f'=IFERROR(ROUND(AVERAGEIFS({PCT_RNG},{SLEEP_RNG},"<7")*100,0)&"%","—")'),
        ("Прогресс при энергии ≥ 4", f'=IFERROR(ROUND(AVERAGEIFS({PCT_RNG},{ENERGY_RNG},">=4")*100,0)&"%","—")'),
    ]
    for j, (label, f) in enumerate(insights):
        r0 = R_CHART0 + 2 * j
        ws.merge_range(r0, C_DONE, r0, C_BEST, label, il)
        ws.merge_range(r0 + 1, C_DONE, r0 + 1, C_BEST, "", iv)
        ctx.fx(ws, r0 + 1, C_DONE, f, iv)
    ws.merge_range(R_CHART0 + 12, C_DONE, R_CHART1, C_BEST, "", ctx.F(bg_color=th["panel"]))

    # --- дневник месяца ---
    dt_f = ctx.F(bg_color=th["bg"], font_color=th["text"], font_size=11, bold=True, valign="vcenter")
    ws.merge_range(R_DT, C_CAT, R_DT, C_D1 + 12, "ДНЕВНИК МЕСЯЦА", dt_f)
    ws.merge_range(R_DT, C_D1 + 13, R_DT, C_BEST,
                   "Одна строка в день: главное событие, урок или победа. Это твоя история месяца.",
                   ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=8, italic=True,
                         align="right", valign="vcenter"))
    ws.write(R_DH, C_CAT, "День", hdr)
    ws.write(R_DH, C_HABIT, "Сделано", hdr)
    ws.merge_range(R_DH, C_TARGET, R_DH, C_BEST, "Итог дня", hdr)
    dl = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, valign="vcenter", align="left", indent=1)
    dc_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=8, bold=True, valign="vcenter", align="left", indent=1)
    dtext = ctx.F(bg_color=th["input"], font_color=th["text"], font_size=9, valign="vcenter", align="left", indent=1)
    for k in range(1, 32):
        r = R_D0 + k - 1
        ctx.fx(ws, r, C_CAT, f'=IF({k}<={DIM},CHOOSE(WEEKDAY({FIRST}+{k - 1},2),{WD_LIST})&" {k} {short}","")', dl)
        ctx.fx(ws, r, C_HABIT, f'=IF(INDEX({CNT_RNG},{k})="","",INDEX({CNT_RNG},{k})&"/"&{NH}&" ✓")', dc_f)
        ws.merge_range(r, C_TARGET, r, C_BEST, mdata["diary"].get(k, ""), dtext)
    b0 = f"${hc}{R_D0 + 1}"
    ws.conditional_format(R_D0, C_CAT, R_D31, C_BEST, {
        "type": "formula", "criteria": f'={b0}=""', "format": ctx.F(bg_color=th["bg"], font_color=th["bg"]),
        "stop_if_true": True})
    ws.conditional_format(R_D0, C_CAT, R_D31, C_BEST, {
        "type": "formula", "criteria": f"=AND({b0}<>\"\",{FIRST}+ROW()-{R_D0 + 1}=TODAY())",
        "format": ctx.F(bg_color=th["today"]), "stop_if_true": True})
    ws.conditional_format(R_D0, C_CAT, R_D31, C_BEST, {
        "type": "formula", "criteria": f'=OR(LEFT({b0},2)="Сб",LEFT({b0},2)="Вс")',
        "format": ctx.F(bg_color=th["weekend"])})

    # --- вид и печать ---
    ws.freeze_panes(0, C_D1)
    ws.set_landscape()
    ws.set_paper(9)
    ws.fit_to_pages(1, 0)
    ws.set_margins(0.3, 0.3, 0.4, 0.4)
    ws.print_area(0, C_MARGIN, R_D31, C_BEST)
    return ws


def add_category_pills(ctx, ws, r1, c1, r2, c2):
    th = ctx.th
    top = f"${cn(c1)}{r1 + 1}"
    for cat, color in CATEGORIES:
        fmt = ctx.F(bg_color=mix(color, th["bg"], th["pill_mix"]),
                    font_color=mix(color, "#000000", th["pill_font"]), bold=True)
        ws.conditional_format(r1, c1, r2, c2, {
            "type": "formula", "criteria": f'={top}="{cat}"', "format": fmt})


# ----------------------------------------------------------------------------
# Лист «Старт»
# ----------------------------------------------------------------------------
def build_start(ctx, habits, year, motto=""):
    th, wb = ctx.th, ctx.wb
    ws = wb.add_worksheet("Старт")
    ws.hide_gridlines(2)
    ws.set_zoom(100)
    ws.set_tab_color(th["accent"])
    ws.set_default_row(16)
    base = ctx.F(bg_color=th["bg"], font_color=th["text"])
    ws.set_column(0, 16383, None, base)
    for c, w in ((0, 1.5), (1, 4.5), (2, 14), (3, 34), (4, 13), (5, 36), (6, 2.5), (7, 82), (8, 1.5)):
        ws.set_column(c, c, w, base)

    ws.set_row(0, 6)
    ws.set_row(1, 34)
    ws.set_row(2, 16)
    ws.set_row(3, 10)
    title = ctx.F(bg_color=th["bg"], font_color=th["text"], bold=True, font_size=24, valign="vcenter")
    ws.merge_range(1, 1, 1, 7, "ПЛАНЕР ПРИВЫЧЕК", title)
    sub = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=10, valign="vcenter")
    ws.merge_range(2, 1, 2, 7,
                   "Один файл на весь год · 12 месяцев · чекбоксы в один клик · статистика, серии и инсайты считаются сами",
                   sub)

    sec = ctx.F(bg_color=th["bg"], font_color=th["accent"], font_size=9, bold=True, valign="vcenter")
    lab = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=9, valign="vcenter", align="left", indent=1)
    inp = ctx.F(bg_color=th["input"], font_color=th["text"], font_size=10, valign="vcenter", align="left",
                indent=1, border=1, border_color=th["line"])
    inp_c = ctx.F(bg_color=th["input"], font_color=th["text"], font_size=10, valign="vcenter", align="center",
                  border=1, border_color=th["line"])
    inp_b = ctx.F(bg_color=th["input"], font_color=th["text"], font_size=10, bold=True, valign="vcenter",
                  align="left", indent=1, border=1, border_color=th["line"])
    note = ctx.F(bg_color=th["bg"], font_color=th["dim"], font_size=8, valign="vcenter", align="left", indent=1)

    # Настройки
    ws.set_row(4, 20)
    ws.merge_range(4, 1, 4, 5, "НАСТРОЙКИ", sec)
    for r in (5, 6, 7):
        ws.set_row(r, 24)
    ws.write(5, 2, "Год", lab)
    ws.write_number(5, 3, year, inp_c)
    ws.data_validation(5, 3, 5, 3, {"validate": "integer", "criteria": "between", "minimum": 2000, "maximum": 2100,
                                    "error_title": "Год", "error_message": "Введи год числом, например 2026"})
    ws.merge_range(5, 4, 5, 5, "все 12 месяцев подстроятся под этот год", note)
    ws.write(6, 2, "Имя", lab)
    ws.write(6, 3, "", inp)
    ws.merge_range(6, 4, 6, 5, "необязательно", note)
    ws.write(7, 2, "Девиз года", lab)
    ws.merge_range(7, 3, 7, 5, motto, inp)

    # Привычки
    ws.set_row(8, 10)
    ws.set_row(9, 24)
    ws.merge_range(9, 1, 9, 5, "МОИ ПРИВЫЧКИ", sec)
    ws.set_row(10, 30)
    hdr = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter",
                align="left", indent=1, text_wrap=True, bottom=1, bottom_color=th["line2"])
    hdr_c = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter",
                  align="center", text_wrap=True, bottom=1, bottom_color=th["line2"])
    ws.write(10, 1, "№", hdr_c)
    ws.write(10, 2, "Сфера", hdr)
    ws.write(10, 3, "Привычка", hdr)
    ws.write(10, 4, "Цель, дней в месяц", hdr_c)
    ws.write(10, 5, "Зачем мне это", hdr)
    ws.write_comment(10, 4, "Сколько дней в месяц планируешь. Пусто = каждый день. "
                            "Например, тренировки 3 раза в неделю ≈ 12.", {"x_scale": 1.5})
    num = ctx.F(bg_color=th["panel"], font_color=th["dim"], font_size=9, align="center", valign="vcenter")
    for i in range(N_HABITS):
        r = 11 + i
        ws.set_row(r, 24)
        ws.write_number(r, 1, i + 1, num)
        h = habits[i] if i < len(habits) else (None, None, None, None)
        cat, name, target, why = h
        ws.write(r, 2, cat or "", inp_c)
        ws.write(r, 3, name or "", inp_b)
        if target:
            ws.write_number(r, 4, target, inp_c)
        else:
            ws.write_blank(r, 4, None, inp_c)
        ws.write(r, 5, why or "", inp)
    ws.data_validation(11, 2, 11 + N_HABITS - 1, 2, {
        "validate": "list", "source": f"={LISTS}!$A$2:$A${1 + len(CATEGORIES)}", "ignore_blank": True,
        "error_title": "Сфера", "error_message": "Выбери сферу из списка (или оставь пустой)"})
    ws.data_validation(11, 4, 11 + N_HABITS - 1, 4, {
        "validate": "integer", "criteria": "between", "minimum": 1, "maximum": 31, "ignore_blank": True,
        "error_title": "Цель", "error_message": "Число дней в месяц от 1 до 31. Пусто = каждый день"})
    add_category_pills(ctx, ws, 11, 2, 11 + N_HABITS - 1, 2)

    r = 11 + N_HABITS
    ws.set_row(r, 8)
    legend = [
        "Ячейки с золотистым фоном — для заполнения. Всё остальное считается автоматически.",
        "Впиши 5–10 привычек: больше — перегруз, меньше — мало данных для статистики.",
        "Пример уже заполнен — просто замени его своими привычками.",
        "Название на конкретном месяце можно перезаписать прямо на листе месяца.",
    ]
    for j, text in enumerate(legend):
        ws.set_row(r + 1 + j, 15)
        ws.merge_range(r + 1 + j, 1, r + 1 + j, 5, text, note)

    # Инструкция справа
    ih = ctx.F(bg_color=th["bg"], font_color=th["accent"], font_size=9, bold=True, valign="vcenter")
    it = ctx.F(bg_color=th["bg"], font_color=th["text"], font_size=9, valign="vcenter", text_wrap=True)
    blocks = [
        ("КАК ПОЛЬЗОВАТЬСЯ", [
            "1. Впиши год и свои привычки слева. Они появятся на всех 12 листах-месяцах.",
            "2. «Цель, дней в месяц» — для привычек не на каждый день (зал 3 раза в неделю ≈ 12). Пусто = ежедневно.",
            "3. Открой лист текущего месяца и каждый вечер ставь галочки. Один клик — готово.",
            "4. Под сеткой отмечай сон, энергию и настроение — планер покажет, как они влияют на результат.",
            "5. В дневнике внизу листа пиши одну строку в день: событие, урок, победа.",
            "6. Раз в неделю смотри «Прогресс», серии 🔥, слабый день и «Подтянуть» — и корректируй план.",
        ]),
        ("ЧТО СЧИТАЕТСЯ САМО", [
            "• Прогресс каждой привычки, каждого дня, недели и месяца — относительно уже прошедших дней.",
            "• Серии подряд 🔥 и личный рекорд по каждой привычке.",
            "• Лучший и слабый день, топ-привычка и что подтянуть, идеальные дни (100%).",
            "• Инсайты: средний сон и энергия, прогресс при хорошем и плохом сне.",
            "• Лист «Год»: итоги по месяцам, график и тепловая карта «привычки × месяцы».",
        ]),
        ("СОВМЕСТИМОСТЬ", [
            "• Excel 365 (Windows, Mac, Web, iPhone/iPad, Android) — чекбоксы работают сразу.",
            "• Google Таблицы: после импорта выдели область галочек → Вставка → Флажок. Формулы работают.",
            "• Старые версии Excel: вместо галочки ставь 1 — всё считается так же.",
            "• Формулы в серых ячейках лучше не трогать. Скрытые колонки справа от статистики — служебные.",
        ]),
    ]
    # строки справа делят высоту со строками таблицы слева, поэтому высоты не трогаем,
    # а строки-разделители (8 и 11+N_HABITS) просто пропускаем
    free_rows = iter(r for r in range(4, 60) if r not in (8, 11 + N_HABITS))
    for head, lines in blocks:
        ws.write(next(free_rows), 7, head, ih)
        for line in lines:
            ws.write(next(free_rows), 7, line, it)
    ws.set_landscape()
    ws.set_paper(9)
    ws.fit_to_pages(1, 1)
    return ws


# ----------------------------------------------------------------------------
# Лист «Год»
# ----------------------------------------------------------------------------
def build_year(ctx):
    th, wb = ctx.th, ctx.wb
    ws = wb.add_worksheet("Год")
    S = "'Год'"
    ws.hide_gridlines(2)
    ws.set_zoom(90)
    ws.set_tab_color("#8B5CF6")
    ws.set_default_row(16)
    base = ctx.F(bg_color=th["bg"], font_color=th["text"])
    ws.set_column(0, 16383, None, base)
    ws.set_column(0, 0, 1.5, base)
    ws.set_column(1, 1, 20, base)
    ws.set_column(2, 14, 9, base)
    ws.set_column(15, 15, 1.5, base)

    ws.set_row(0, 6)
    ws.set_row(1, 34)
    ws.set_row(2, 16)
    ws.set_row(3, 10)
    title = ctx.F(bg_color=th["bg"], font_color=th["text"], bold=True, font_size=24, valign="vcenter")
    ws.merge_range(1, 1, 1, 14, "", title)
    ctx.fx(ws, 1, 1, f'="ГОД "&{YEAR_REF}&" — ИТОГИ"', title)
    ws.merge_range(2, 1, 2, 14, "Месяцы без отметок остаются пустыми. Всё обновляется само.",
                   ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=10, valign="vcenter"))

    # KPI
    R_KL_, R_KV_, R_KS_ = 4, 5, 6
    ws.set_row(R_KL_, 14)
    ws.set_row(R_KV_, 28)
    ws.set_row(R_KS_, 14)
    ws.set_row(7, 10)
    RT, R0 = 8, 9                       # заголовок таблицы, первая строка месяцев
    RTOT = R0 + 12
    tile_kw = dict(bg_color=th["tile"], left=5, left_color=th["accent"], align="left", indent=1)
    kl = ctx.F(font_color=th["muted"], font_size=8, bold=True, valign="bottom", **tile_kw)
    kv = ctx.F(font_color=th["text"], font_size=18, bold=True, valign="vcenter", **tile_kw)
    kv_pct = ctx.F(font_color=th["accent"], font_size=18, bold=True, valign="vcenter", num_format="0%", **tile_kw)
    kv_days = ctx.F(font_color=th["text"], font_size=18, bold=True, valign="vcenter", num_format='0" дн."', **tile_kw)
    ks = ctx.F(font_color=th["muted"], font_size=8, valign="top", **tile_kw)
    prog_col = R(R0, 2, R0 + 11, 2)
    names_col = R(R0, 1, R0 + 11, 1)
    tiles = [
        ("ПРОГРЕСС ГОДА", f"={A(RTOT, 2, True, True)}", kv_pct, "средний прогресс начатых месяцев"),
        ("ОТМЕТОК ЗА ГОД", f"={A(RTOT, 3, True, True)}", kv, "выполненных привычек"),
        ("ИДЕАЛЬНЫХ ДНЕЙ", f"={A(RTOT, 4, True, True)}", kv, "дней, когда сделано всё"),
        ("ЛУЧШИЙ МЕСЯЦ", f'=IFERROR(INDEX({names_col},MATCH(MAX({prog_col}),{prog_col},0)),"—")', kv,
         "по прогрессу"),
        ("РЕКОРД СЕРИИ", f"={A(RTOT, 5, True, True)}", kv_days, "самая длинная серия года"),
    ]
    spans = [(1, 2), (3, 5), (6, 8), (9, 11), (12, 14)]
    for (label, fval, vfmt, subtext), (c1, c2) in zip(tiles, spans):
        ws.merge_range(R_KL_, c1, R_KL_, c2, label, kl)
        ws.merge_range(R_KV_, c1, R_KV_, c2, "", vfmt)
        ctx.fx(ws, R_KV_, c1, fval, vfmt)
        ws.merge_range(R_KS_, c1, R_KS_, c2, subtext, ks)

    # Таблица по месяцам
    hdr = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter",
                align="center", text_wrap=True, bottom=1, bottom_color=th["line2"])
    hdr_l = ctx.F(bg_color=th["panel"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter",
                  align="left", indent=1, bottom=1, bottom_color=th["line2"])
    ws.set_row(RT, 30)
    heads = ["Месяц", "Прогресс", "Отметок", "Идеальных дней", "Лучшая серия", "Сон, ч", "Энергия"]
    for j, h in enumerate(heads):
        ws.write(RT, 1 + j, h, hdr_l if j == 0 else hdr)
    m_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=10, bold=True, valign="vcenter",
                align="left", indent=1)
    v_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=10, valign="vcenter", align="center")
    p_f = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=10, valign="vcenter", align="center",
                num_format="0%")
    for mi in range(1, 13):
        r = R0 + mi - 1
        ws.set_row(r, 20)
        MS = f"'{MONTHS[mi - 1][0]}'"
        off = f"OR({MS}!{EL}=0,{MS}!{DONE}=0)"      # месяц не начался или без единой отметки
        ws.write(r, 1, MONTHS[mi - 1][0], m_f)
        ctx.fx(ws, r, 2, f'=IF({off},"",{MS}!{PROG})', p_f)
        ctx.fx(ws, r, 3, f'=IF({off},"",{MS}!{DONE})', v_f)
        ctx.fx(ws, r, 4, f'=IF({off},"",{MS}!{PERF})', v_f)
        ctx.fx(ws, r, 5, f'=IF({off},"",MAX({MS}!{BEST_RNG}))', v_f)
        ctx.fx(ws, r, 6, f'=IF({off},"",IFERROR(ROUND(AVERAGE({MS}!{SLEEP_RNG}),1),"—"))', v_f)
        ctx.fx(ws, r, 7, f'=IF({off},"",IFERROR(ROUND(AVERAGE({MS}!{ENERGY_RNG}),1),"—"))', v_f)
    ws.set_row(RTOT, 22)
    t_l = ctx.F(bg_color=th["panel2"], font_color=th["text"], font_size=10, bold=True, valign="vcenter",
                align="left", indent=1, top=1, top_color=th["line2"])
    t_v = ctx.F(bg_color=th["panel2"], font_color=th["accent"], font_size=10, bold=True, valign="vcenter",
                align="center", top=1, top_color=th["line2"])
    t_p = ctx.F(bg_color=th["panel2"], font_color=th["accent"], font_size=10, bold=True, valign="vcenter",
                align="center", num_format="0%", top=1, top_color=th["line2"])
    ws.write(RTOT, 1, "Итого / среднее", t_l)
    ctx.fx(ws, RTOT, 2, f"=IFERROR(AVERAGE({prog_col}),0)", t_p)
    ctx.fx(ws, RTOT, 3, f"=SUM({R(R0, 3, R0 + 11, 3)})", t_v)
    ctx.fx(ws, RTOT, 4, f"=SUM({R(R0, 4, R0 + 11, 4)})", t_v)
    ctx.fx(ws, RTOT, 5, f"=MAX(0,MAX({R(R0, 5, R0 + 11, 5)}))", t_v)
    ctx.fx(ws, RTOT, 6, f'=IFERROR(ROUND(AVERAGE({R(R0, 6, R0 + 11, 6)}),1),"—")', t_v)
    ctx.fx(ws, RTOT, 7, f'=IFERROR(ROUND(AVERAGE({R(R0, 7, R0 + 11, 7)}),1),"—")', t_v)
    ws.conditional_format(R0, 2, R0 + 11, 2, {
        "type": "data_bar", "bar_color": th["green"], "bar_solid": True, "bar_no_border": True,
        "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 1, "bar_axis_position": "none"})
    ws.conditional_format(R0, 1, R0 + 11, 7, {
        "type": "formula", "criteria": f"=AND(YEAR(TODAY())={YEAR_REF},MONTH(TODAY())=ROW()-{R0})",
        "format": ctx.F(bg_color=th["today"])})

    # График по месяцам
    RC = RTOT + 2
    sec = ctx.F(bg_color=th["bg"], font_color=th["muted"], font_size=8, bold=True, valign="vcenter")
    ws.set_row(RC, 18)
    ws.merge_range(RC, 1, RC, 14, "ПРОГРЕСС ПО МЕСЯЦАМ", sec)
    chart = wb.add_chart({"type": "column"})
    chart.add_series({
        "name": "Прогресс", "categories": f"={S}!{names_col}", "values": f"={S}!{prog_col}",
        "fill": {"color": th["accent"]}, "border": {"none": True}, "gap": 60,
        "data_labels": {"value": True, "num_format": "0%", "font": {"name": FONT, "color": th["muted"], "size": 8}},
    })
    chart.set_legend({"none": True})
    chart.set_title({"none": True})
    chart.set_chartarea({"fill": {"color": th["chart_bg"]}, "border": {"none": True}})
    chart.set_plotarea({"fill": {"none": True}})
    axis_font = {"name": FONT, "color": th["muted"], "size": 8}
    chart.set_x_axis({"num_font": axis_font, "line": {"color": th["grid"]}, "major_gridlines": {"visible": False},
                      "major_tick_mark": "none"})
    chart.set_y_axis({"min": 0, "max": 1, "major_unit": 0.25, "num_format": "0%", "num_font": axis_font,
                      "line": {"none": True}, "major_gridlines": {"visible": True, "line": {"color": th["grid"]}}})
    chart.set_size({"width": 20 * 7 + 5 + 13 * (9 * 7 + 5), "height": 14 * 20})
    ws.insert_chart(RC + 1, 1, chart)
    for r in range(RC + 1, RC + 15):
        ws.set_row(r, 15)

    # Тепловая карта
    RH = RC + 16
    ws.set_row(RH, 18)
    ws.merge_range(RH, 1, RH, 14, "КАРТА ПРИВЫЧЕК: ПРОГРЕСС ПО МЕСЯЦАМ", sec)
    ws.set_row(RH + 1, 20)
    ws.write(RH + 1, 1, "Привычка", hdr_l)
    for mi in range(1, 13):
        ws.write(RH + 1, 1 + mi, MONTHS[mi - 1][1], hdr)
    ws.write(RH + 1, 14, "Год", hdr)
    hm_n = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=9, bold=True, valign="vcenter",
                 align="left", indent=1)
    hm_v = ctx.F(bg_color=th["panel"], font_color=th["text"], font_size=8, valign="vcenter", align="center",
                 num_format="0%", border=1, border_color=th["bg"])
    for i in range(N_HABITS):
        r = RH + 2 + i
        ws.set_row(r, 18)
        sr = R_H0 + i + 1
        ctx.fx(ws, r, 1, f"=IF('Старт'!$D${sr}=\"\",\"\",'Старт'!$D${sr})", hm_n)
        for mi in range(1, 13):
            MS = f"'{MONTHS[mi - 1][0]}'"
            cell = f"{MS}!{A(R_H0 + i, C_PROG, True, True)}"
            ctx.fx(ws, r, 1 + mi, f'=IF(OR({MS}!{EL}=0,{MS}!{DONE}=0,{cell}=""),"",{cell})', hm_v)
        ctx.fx(ws, r, 14, f'=IFERROR(AVERAGE({R(r, 2, r, 13)}),"")', hm_v)
    ws.conditional_format(RH + 2, 2, RH + 1 + N_HABITS, 14, {
        "type": "3_color_scale", "min_type": "num", "min_value": 0, "mid_type": "num", "mid_value": 0.5,
        "max_type": "num", "max_value": 1, "min_color": th["heat"][0], "mid_color": th["heat"][1],
        "max_color": th["heat"][2]})
    ws.set_landscape()
    ws.set_paper(9)
    ws.fit_to_pages(1, 0)
    ws.print_area(0, 0, RH + 1 + N_HABITS, 15)
    return ws


# ----------------------------------------------------------------------------
# Скрытые справочники
# ----------------------------------------------------------------------------
def build_lists(ctx):
    ws = ctx.wb.add_worksheet("Списки")
    ws.write(0, 0, "Сферы")
    for i, (cat, _) in enumerate(CATEGORIES):
        ws.write(1 + i, 0, cat)
    ws.write(0, 1, "Энергия")
    for i in range(5):
        ws.write_number(1 + i, 1, i + 1)
    ws.write(0, 2, "Настроение")
    for i, m in enumerate(MOODS):
        ws.write(1 + i, 2, m)
    ws.write(0, 3, "Месяцы")
    for i, (m, _) in enumerate(MONTHS):
        ws.write(1 + i, 3, m)
    ws.hide()
    return ws


# ----------------------------------------------------------------------------
# Сборка книги
# ----------------------------------------------------------------------------
def build_workbook(path, theme, data, year, today, cache=None, preview=False, habits=None, motto=""):
    wb = xlsxwriter.Workbook(path, {"default_format_properties": {"font_name": FONT, "font_size": 10}})
    ctx = Ctx(wb, THEMES[theme], cache, preview)
    build_start(ctx, habits or EXAMPLE_HABITS, year, motto)
    build_year(ctx)
    months = [build_month(ctx, mi, data.get(mi, EMPTY), today) for mi in range(1, 13)]
    build_lists(ctx)
    active = months[today.month - 1] if today.year == year else months[0]
    active.activate()
    wb.set_properties({"title": "Планер привычек", "subject": "Трекер привычек на год",
                       "comments": "Собрано скриптом build_planner.py"})
    wb.close()
    return ctx.formula_cells


LO_RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def lo_recalc(path, cells, timeout=600):
    """Пересчитать формулы через LibreOffice (headless, с принудительным пересчётом при загрузке).

    Возвращает (cache, errors): cache — значения формул по адресам, errors — ячейки с ошибками.
    """
    import tempfile
    import warnings
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(prefix="planner_lo_") as tmp:
        profile = os.path.join(tmp, "profile")
        os.makedirs(os.path.join(profile, "user"))
        with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w", encoding="utf-8") as f:
            f.write(LO_RECALC_XCU)
        out = os.path.join(tmp, "out")
        os.makedirs(out)
        env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
        cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
               "--convert-to", "xlsx", "--outdir", out, path]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=env)
        result = os.path.join(out, os.path.basename(path))
        if not os.path.exists(result):
            raise RuntimeError("LibreOffice не пересчитал файл: " + (res.stdout + res.stderr)[-400:])
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(result, data_only=True)
    cache, errors = {}, []
    epoch = dt.datetime(1899, 12, 30)
    for sheet, row, col in cells:
        v = wb[sheet].cell(row + 1, col + 1).value
        if isinstance(v, dt.datetime):
            v = (v - epoch).days + (v - epoch).seconds / 86400
        elif isinstance(v, dt.date):
            v = (dt.datetime(v.year, v.month, v.day) - epoch).days
        elif isinstance(v, str) and v.startswith("#"):
            errors.append(f"{sheet}!{rc(row, col)} = {v}")
        cache[(sheet, row, col)] = v
    return cache, errors


def build(path, theme, data, year, today, do_recalc=False, preview=False, habits=None, motto=""):
    cells = build_workbook(path, theme, data, year, today, None, preview, habits, motto)
    if not do_recalc:
        return None
    cache, errors = lo_recalc(path, cells)
    print(f"  пересчёт LibreOffice: формул {len(cells)}, ошибок {len(errors)}")
    if errors:
        print("   " + "\n   ".join(errors[:40]))
        raise SystemExit("формулы содержат ошибки — файл не собран")
    build_workbook(path, theme, data, year, today, cache, preview, habits, motto)
    return cache


# ----------------------------------------------------------------------------
# Демо-данные (для файла-примера и проверки формул)
# ----------------------------------------------------------------------------
def make_demo(year, today):
    rnd = random.Random(7)
    probs = [0.93, 0.5, 0.72, 0.8, 0.6, 0.55, 0.85, 0.66]      # по числу примерных привычек
    data = {}
    months = [m for m in (today.month - 2, today.month - 1, today.month) if 1 <= m <= 12]
    for mi in months:
        dim = calendar.monthrange(year, mi)[1]
        last = dim if mi < today.month else today.day
        boost = {months[0]: -0.08, months[-1]: 0.05}.get(mi, 0.0)
        d = dict(checks=set(), sleep={}, energy={}, mood={}, diary={})
        for day in range(1, last + 1):
            sleep = rnd.choice([5.5, 6, 6.5, 7, 7, 7.5, 7.5, 8, 8.5])
            factor = 1.15 if sleep >= 7 else 0.72
            done = 0
            for h, p in enumerate(probs):
                if rnd.random() < min(0.97, (p + boost) * factor):
                    d["checks"].add((h, day))
                    done += 1
            ratio = done / len(probs)
            energy = max(1, min(5, round(1 + 4 * ratio + rnd.uniform(-0.8, 0.8))))
            mood = MOODS[0] if ratio > 0.85 else MOODS[1] if ratio > 0.65 else MOODS[2] if ratio > 0.45 \
                else MOODS[3] if ratio > 0.25 else MOODS[4]
            d["sleep"][day], d["energy"][day], d["mood"][day] = sleep, energy, mood
        data[mi] = d
    cur = data[today.month]
    perfect = min(12, today.day)
    for h in range(len(probs)):                      # один гарантированно идеальный день
        cur["checks"].add((h, perfect))
    cur["sleep"][perfect], cur["energy"][perfect], cur["mood"][perfect] = 8, 5, MOODS[0]
    for h in range(len(probs)):                      # длинная серия у первой привычки
        pass
    for day in range(1, today.day + 1):
        cur["checks"].add((0, day)) if day != 3 else cur["checks"].discard((0, day))
    cur["diary"].update({
        1: "Начал месяц с плана: 8 привычек, фокус — сон и спорт.",
        5: "Пропустил зал, зато прошёл 12 000 шагов и лёг вовремя.",
        perfect: "Идеальный день: всё выполнено. Рано лёг — легко встал.",
        min(today.day, 18): "Урок: если не спланировать вечер, чтение выпадает первым.",
    })
    return data


def expected_stats(data, mi, year, today, n_named):
    """Ожидаемые значения для проверки формул (текущий месяц)."""
    d = data[mi]
    dim = calendar.monthrange(year, mi)[1]
    el = today.day if (today.year == year and today.month == mi) else (dim if dt.date(year, mi, 1) < today else 0)
    per_day = {k: sum(1 for h in range(n_named) if (h, k) in d["checks"]) for k in range(1, el + 1)}
    per_habit, streak, best = {}, {}, {}
    for h in range(n_named):
        per_habit[h] = sum(1 for k in range(1, dim + 1) if (h, k) in d["checks"])
        run, runs = 0, [0] * (el + 1)
        for k in range(1, el + 1):
            run = run + 1 if (h, k) in d["checks"] else 0
            runs[k] = run
        best[h] = max(runs)
        streak[h] = max(runs[el], runs[el - 1] if el > 1 else 0) if el else 0
    perfect = sum(1 for k, v in per_day.items() if v == n_named)
    return dict(el=el, dim=dim, per_day=per_day, per_habit=per_habit, streak=streak, best=best, perfect=perfect)


def verify(cache, data, mi, year, today, n_named):
    """Сравнить кэш LibreOffice с независимым расчётом на Python."""
    exp = expected_stats(data, mi, year, today, n_named)
    sheet = MONTHS[mi - 1][0]
    g = lambda r, c: cache.get((sheet, r, c))
    problems = []
    checks = [("дней прошло", g(2, C_H), exp["el"]), ("дней в месяце", g(1, C_H), exp["dim"]),
              ("привычек", g(3, C_H), n_named), ("идеальных дней", g(8, C_H), exp["perfect"])]
    for k, v in exp["per_day"].items():
        checks.append((f"день {k}: сделано", g(R_CNT, C_D1 + k - 1), v))
    for h in range(n_named):
        checks += [(f"привычка {h + 1}: сделано", g(R_H0 + h, C_DONE), exp["per_habit"][h]),
                   (f"привычка {h + 1}: серия", g(R_H0 + h, C_STREAK), exp["streak"][h]),
                   (f"привычка {h + 1}: рекорд", g(R_H0 + h, C_BEST), exp["best"][h])]
    for label, got, want in checks:
        if got != want:
            problems.append(f"{label}: в файле {got!r}, ожидалось {want!r}")
    return problems


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Сборка планера привычек")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "dist"))
    ap.add_argument("--recalc", action="store_true",
                    help="пересчитать формулы в LibreOffice и записать значения в файл (нужен soffice)")
    ap.add_argument("--preview", action="store_true", help="собрать также вариант для рендера превью")
    ap.add_argument("--year", type=int, default=None)
    args = ap.parse_args()

    today = dt.date.today()
    year = args.year or today.year
    os.makedirs(args.out, exist_ok=True)
    demo = make_demo(year, today)

    targets = [
        ("Планер привычек.xlsx", "dark", {}, False),
        ("Планер привычек (светлая тема).xlsx", "light", {}, False),
        ("Планер привычек (пример с данными).xlsx", "dark", demo, False),
    ]
    if args.preview:
        targets += [("preview_dark.xlsx", "dark", demo, True), ("preview_light.xlsx", "light", demo, True)]
    for fname, theme, data, preview in targets:
        path = os.path.join(args.out, fname)
        print(f"→ {fname}")
        motto = "Не идеально, а регулярно." if data else ""
        cache = build(path, theme, data, year, today, args.recalc, preview, motto=motto)
        if cache is not None and data:
            problems = verify(cache, data, today.month, year, today, len(EXAMPLE_HABITS))
            if problems:
                print("  ПРОВЕРКА НЕ ПРОШЛА:\n   " + "\n   ".join(problems[:40]))
                raise SystemExit(1)
            print(f"  проверка формул по демо-данным: ок ({len(cache)} значений)")
    print("готово:", args.out)


if __name__ == "__main__":
    main()
