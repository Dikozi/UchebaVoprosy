#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Эталонная модель «Режима»: считает те же величины, что и формулы книги,
но независимо, на Python. Расхождение — красная сборка (раздел 8 спецификации).

    python3 model.py <пересчитанный.xlsx> <демо.json> --year 2026 --today 2026-09-21
"""
import argparse
import calendar
import datetime as dt
import json
import sys
import warnings

warnings.filterwarnings("ignore")
from openpyxl import load_workbook                                   # noqa: E402

sys.path.insert(0, __file__.rsplit("/", 1)[0])
import texts as T                                                    # noqa: E402
import build as B                                                    # noqa: E402


def simulate(demo, year, today, freqs, active):
    """Повторяет движок «Расчёта»: серии со щитом, опыт, полные дни, возвраты."""
    total = 366 if calendar.isleap(year) else 365
    nh = len(freqs)
    streak = [0] * nh
    miss = [0] * nh
    xp = 0
    st = {"ticks": 0, "full": 0, "returns": 0, "record": 0, "evenings": 0}
    per_month = {m: {"done": 0, "full": 0, "xp": 0, "best": 0, "el": 0} for m in range(1, 13)}
    per_habit_month = {(m, i): 0 for m in range(1, 13) for i in range(nh)}
    streak_today = [0] * nh
    miss_today = [0] * nh
    eve_run = 0
    for yday in range(1, total + 1):
        date = dt.date(year, 1, 1) + dt.timedelta(days=yday - 1)
        passed = date <= today
        m, d = date.month, date.day
        marks = demo.get(m, {})
        checks = marks.get("checks", set())
        if passed:
            per_month[m]["el"] += 1
        day_done = 0
        day_xp = 0
        for i in range(nh):
            done = (i, d) in checks and active[i]
            if not passed:
                continue
            if done:
                day_done += 1
                st["ticks"] += 1
                per_month[m]["done"] += 1
                per_habit_month[(m, i)] += 1
            prev_s, prev_m = streak[i], miss[i]
            if d == 1:
                prev_m = 0
                miss[i] = 0
            miss[i] = prev_m + (0 if done else 1)
            if done:
                if freqs[i] < 7:
                    day_xp += 10
                elif prev_s == 0 and prev_m > 1:
                    day_xp += 15
                    st["returns"] += 1
                elif prev_s >= 20:
                    day_xp += 15
                elif prev_s >= 6:
                    day_xp += 13
                else:
                    day_xp += 10
                streak[i] = prev_s + 1
            else:
                streak[i] = prev_s if miss[i] <= 1 else 0
            st["record"] = max(st["record"], streak[i])
            per_month[m]["best"] = max(per_month[m]["best"], streak[i])
        if passed:
            daily = [i for i in range(nh) if active[i] and freqs[i] == 7]
            if daily and all((i, d) in checks for i in daily):
                st["full"] += 1
                per_month[m]["full"] += 1
            xp += day_xp
            per_month[m]["xp"] += day_xp
            eve_run = eve_run + 1 if marks.get("sleep", {}).get(d) else 0
            st["evenings"] = max(st["evenings"], eve_run)
        if date == today:
            streak_today = list(streak)
            miss_today = list(miss)
    # испытание недели: слабейшая привычка прошлой недели, планка «5 из 7»
    weeks = B.week_blocks(year)
    per_week = []
    for days in weeks:
        counts = []
        for i in range(nh):
            if not active[i]:
                counts.append(99)
                continue
            c = 0
            for yday in days:
                date = dt.date(year, 1, 1) + dt.timedelta(days=yday - 1)
                if date <= today and (i, date.day) in demo.get(date.month, {}).get("checks", set()):
                    c += 1
            counts.append(c)
        per_week.append(counts)
    closed = 0
    for i, counts in enumerate(per_week):
        if i == 0:
            continue
        prev = per_week[i - 1]
        if min(prev) == 99:
            continue
        weak = prev.index(min(prev))
        if counts[weak] >= 5:
            closed += 1
    st["challenges"] = closed
    xp += 100 * closed
    return dict(xp=xp, streak=streak_today, miss=miss_today, per_month=per_month,
                per_habit_month=per_habit_month, **st)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("demo")
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--today", required=True)
    args = ap.parse_args()

    today = dt.date.fromisoformat(args.today)
    raw = json.load(open(args.demo))
    demo = {int(m): {"checks": {(int(a), int(b)) for a, b in d["checks"]},
                     "sleep": {int(k): v for k, v in d["sleep"].items()}}
            for m, d in raw.items()}
    freqs = [e[3] for e in T.EXAMPLES] + [7] * (B.NH - len(T.EXAMPLES))
    active = [True] * len(T.EXAMPLES) + [False] * (B.NH - len(T.EXAMPLES))
    exp = simulate(demo, args.year, today, freqs, active)

    wb = load_workbook(args.workbook, data_only=True)
    calc = wb[B.CALC]
    rowt = B.kr(today.timetuple().tm_yday) + 1
    problems = []

    def cmp(label, got, want):
        if got != want:
            problems.append(f"{label}: в книге {got!r}, ожидалось {want!r}")

    cmp("строка сегодня", calc["B2"].value, rowt)
    cmp("активных привычек", calc["B3"].value, sum(active))
    cmp("опыт", calc["B4"].value, exp["xp"])
    for i in range(B.NH):
        if not active[i]:
            continue
        cmp(f"серия {i + 1}", calc.cell(rowt, B.KC_S0 + i + 1).value, exp["streak"][i])
        cmp(f"пропуски {i + 1}", calc.cell(rowt, B.KC_MISS0 + i + 1).value, exp["miss"][i])
    for m in range(1, 13):
        r = B.K_MON + m
        cmp(f"{T.MONTHS_SHORT[m - 1]}: галочек", calc.cell(r, B.MC_DONE + 1).value,
            exp["per_month"][m]["done"])
        cmp(f"{T.MONTHS_SHORT[m - 1]}: полных дней", calc.cell(r, B.MC_FULL + 1).value,
            exp["per_month"][m]["full"])
        cmp(f"{T.MONTHS_SHORT[m - 1]}: опыт", calc.cell(r, B.MC_XP + 1).value,
            exp["per_month"][m]["xp"])
        cmp(f"{T.MONTHS_SHORT[m - 1]}: прошло дней", calc.cell(r, B.MC_EL + 1).value,
            exp["per_month"][m]["el"])
    counters = {"ticks": exp["ticks"], "record": exp["record"], "full_days": exp["full"],
                "returns": exp["returns"], "evenings": exp["evenings"],
                "challenges": exp["challenges"]}
    order = ["ticks", "record", "full_days", "returns", "challenges", "evenings", "answers"]
    for i, key in enumerate(order):
        if key in counters:
            cmp(f"счётчик {key}", calc.cell(B.K_AWARD + 1 + i, 2).value, counters[key])

    if problems:
        print("ПРОВЕРКА НЕ ПРОШЛА:\n  " + "\n  ".join(problems[:40]))
        return 1
    print(f"эталонная модель сошлась: опыт {exp['xp']}, галочек {exp['ticks']}, "
          f"рекорд {exp['record']}, полных дней {exp['full']}, возвратов {exp['returns']}, "
          f"испытаний {exp['challenges']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
